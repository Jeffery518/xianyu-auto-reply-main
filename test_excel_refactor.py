import io
import openpyxl
from openpyxl.styles import PatternFill

# Mock data
keywords = [
    {'keyword': '你好', 'item_id': None, 'reply': '您好！', 'type': 'text'},
    {'keyword': '价格', 'item_id': '123', 'reply': '99元', 'type': 'text'},
    {'keyword': '图片', 'item_id': None, 'reply': 'img.jpg', 'type': 'image'}
]

def mock_export(keywords):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "关键词数据"
    headers = ['关键词', '商品ID', '关键词内容']
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)
    data_rows = []
    for keyword_data in keywords:
        if keyword_data.get('type', 'text') == 'text':
            data_rows.append([
                keyword_data['keyword'],
                keyword_data['item_id'] or '',
                keyword_data['reply']
            ])
    if data_rows:
        for row_idx, row_data in enumerate(data_rows, 2):
            for col_idx, value in enumerate(row_data, 1):
                ws.cell(row=row_idx, column=col_idx, value=value)
    else:
        # Template logic
        pass
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def mock_import(file_contents):
    wb = openpyxl.load_workbook(io.BytesIO(file_contents), data_only=True)
    ws = wb.active
    headers = []
    for cell in ws[1]:
        val = str(cell.value).strip() if cell.value is not None else ""
        headers.append(val)
    required_columns = ['关键词', '商品ID', '关键词内容']
    col_map = {}
    for req in required_columns:
        if req in headers:
            col_map[req] = headers.index(req) + 1
        else:
            raise Exception(f"Missing column: {req}")
    import_data = []
    for row_idx in range(2, ws.max_row + 1):
        keyword_cell = ws.cell(row=row_idx, column=col_map['关键词']).value
        item_id_cell = ws.cell(row=row_idx, column=col_map['商品ID']).value
        reply_cell = ws.cell(row=row_idx, column=col_map['关键词内容']).value
        if keyword_cell is None: continue
        keyword = str(keyword_cell).strip()
        item_id = str(item_id_cell).strip() if item_id_cell is not None and str(item_id_cell).strip() else None
        reply = str(reply_cell).strip() if reply_cell is not None else ""
        if not keyword: continue
        import_data.append((keyword, reply, item_id))
    return import_data

# Test execution
print("Testing Export...")
exported_file = mock_export(keywords)
print(f"Exported size: {len(exported_file.getvalue())} bytes")

print("Testing Import...")
imported_data = mock_import(exported_file.getvalue())
print(f"Imported data: {imported_data}")

# Verification
assert len(imported_data) == 2
assert imported_data[0] == ('你好', '您好！', None)
assert imported_data[1] == ('价格', '99元', '123')
print("Verification Success!")
